#!/usr/bin/env python3
"""Download PDFs listed in the Downloader Master tab of an Excel workbook.

Install: py -m pip install openpyxl
Run: py download_icai_pdfs.py --excel "inventory.xlsx" --output "D:\\Local_CACSCMA_RAG_Master"
Or run without arguments and enter both paths when prompted.
Export the updated Google Sheet as Microsoft Excel (.xlsx) first.
Python 3.10+; only third-party dependency is openpyxl. No Google login required.
"""
import argparse
import csv
import os
from pathlib import Path
import re
import sys
import time
from urllib.error import HTTPError, URLError
from urllib.parse import urlsplit
from urllib.request import Request, urlopen


def is_pdf(path):
    try:
        with path.open('rb') as stream:
            return b'%PDF-' in stream.read(1024) and path.stat().st_size > 8
    except OSError:
        return False


def enabled(value):
    return value is True or str(value).strip().lower() in {'true', '1', '1.0', 'yes'}


def destination(root, folder, filename):
    # Validate Windows paths even when this script is tested on another OS.
    parts = str(folder).replace('\\', '/').split('/')
    name = str(filename)
    for part in parts + [name]:
        if (not part or part in {'.', '..'} or re.search(r'[<>:"/\\|?*\x00-\x1f]', part)
                or part.endswith((' ', '.'))
                or re.fullmatch(r'(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?', part)):
            raise ValueError('Unsafe folder or filename: ' + part)
    if not name.lower().endswith('.pdf'):
        raise ValueError('Filename must end in .pdf')
    target = root.joinpath(*parts, name).resolve()
    if not target.is_relative_to(root):
        raise ValueError('Destination escapes the output folder')
    return target


def read_plan(workbook):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Install the dependency first: py -m pip install openpyxl') from exc
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        if 'Downloader Master' not in book.sheetnames:
            raise ValueError('Missing Downloader Master tab. Download the updated Google Sheet as .xlsx.')
        rows = book['Downloader Master'].iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else '' for v in next(rows)]
        required = ['download_enabled', 'download_url', 'relative_folder', 'filename']
        if any(headers.count(h) != 1 for h in required):
            raise ValueError('Expected exactly one of each header: ' + ', '.join(required))
        indexes = [headers.index(h) for h in required]
        return [(row_number, dict(zip(required, [r[i] if i < len(r) else None for i in indexes])))
                for row_number, r in enumerate(rows, 2) if any(v is not None for v in r)]
    finally:
        book.close()


def download(url, target, attempts, timeout):
    parsed = urlsplit(url)
    if parsed.scheme not in {'https', 'http'} or not parsed.hostname or parsed.username:
        raise ValueError('Download URL must be an HTTP(S) URL without embedded credentials')
    target.parent.mkdir(parents=True, exist_ok=True)
    partial = target.with_name(target.name + '.part')
    for attempt in range(attempts):
        try:
            request = Request(url, headers={'User-Agent': 'ICAI-PDF-Downloader/1.0', 'Accept': 'application/pdf'})
            with urlopen(request, timeout=timeout) as response, partial.open('wb') as stream:
                status = response.status
                if status != 200:
                    raise ValueError(f'Unexpected HTTP status {status}')
                total = 0
                while chunk := response.read(1024 * 128):
                    stream.write(chunk)
                    total += len(chunk)
                expected = response.headers.get('Content-Length')
                if expected and total != int(expected):
                    raise ValueError('Incomplete response body')
            if not is_pdf(partial):
                raise ValueError('Response is not a PDF (may be an HTML error page)')
            os.replace(partial, target)
            return status
        except (OSError, URLError, ValueError) as exc:
            partial.unlink(missing_ok=True)
            if isinstance(exc, HTTPError) and exc.code not in {408, 429, 500, 502, 503, 504}:
                raise
            if attempt + 1 == attempts:
                raise
            time.sleep(min(2 ** attempt, 8))
        except KeyboardInterrupt:
            partial.unlink(missing_ok=True)
            raise


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--excel', help='Downloaded .xlsx workbook')
    parser.add_argument('--output', help='Destination root folder')
    parser.add_argument('--attempts', type=int, default=3, help='Maximum attempts per URL (default 3)')
    parser.add_argument('--timeout', type=float, default=60, help='Network timeout in seconds (default 60)')
    parser.add_argument('--delay', type=float, default=0.5, help='Seconds between downloads (default 0.5)')
    parser.add_argument('--dry-run', action='store_true', help='Validate and list destinations without downloading')
    args = parser.parse_args(argv)
    if args.attempts < 1 or args.timeout <= 0 or args.delay < 0:
        parser.error('attempts must be >= 1, timeout > 0, and delay >= 0')
    try:
        excel = Path((args.excel or input('Excel workbook path: ')).strip().strip('"')).expanduser()
        output = (args.output or input('Destination folder: ')).strip().strip('"')
        if not output:
            raise ValueError('Destination folder cannot be empty')
        root = Path(output).expanduser().resolve()
        plan = read_plan(excel)
        # Preflight all enabled rows before making any downloads.
        seen = set()
        targets = {}
        for row_number, row in plan:
            if not enabled(row['download_enabled']):
                continue
            if not all(row[k] for k in ['download_url', 'relative_folder', 'filename']):
                raise ValueError(f'Row {row_number}: missing required value')
            target = destination(root, row['relative_folder'], row['filename'])
            key = str(target).casefold()
            if key in seen:
                raise ValueError(f'Row {row_number}: duplicate destination {target}')
            seen.add(key)
            targets[row_number] = target
        if args.dry_run:
            for target in targets.values():
                print(target)
            print(f'Validated {len(targets)} enabled rows; no files downloaded.')
            return 0
        root.mkdir(parents=True, exist_ok=True)
        log = root / 'download_log.csv'
        counts = dict(downloaded=0, already_present=0, skipped=0, failed=0)
        with log.open('a', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            if stream.tell() == 0:
                writer.writerow(['timestamp_utc', 'excel_row', 'url', 'path', 'outcome', 'http_status', 'error'])
            for number, (row_number, row) in enumerate(plan, 1):
                target = targets.get(row_number)
                status, error = '', ''
                if target is None:
                    outcome = 'skipped'
                elif is_pdf(target):
                    outcome = 'already_present'
                else:
                    try:
                        status = download(str(row['download_url']).strip(), target, args.attempts, args.timeout)
                        outcome = 'downloaded'
                    except (OSError, URLError, ValueError) as exc:
                        outcome, error = 'failed', str(exc)
                        status = getattr(exc, 'code', '')
                    time.sleep(args.delay)
                counts[outcome] += 1
                writer.writerow([time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()), row_number,
                                 row['download_url'], str(target or ''), outcome, status, error])
                stream.flush()
                print(f'[{number}/{len(plan)}] {outcome}: {row["filename"]}' + (f' | {error}' if error else ''))
        print('\n' + ', '.join(f'{k}: {v}' for k, v in counts.items()))
        print(f'Log: {log}')
        return 1 if counts['failed'] else 0
    except KeyboardInterrupt:
        print('\nStopped. Run the same command again to resume.', file=sys.stderr)
        return 130
    except (OSError, ValueError, EOFError, StopIteration) as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
