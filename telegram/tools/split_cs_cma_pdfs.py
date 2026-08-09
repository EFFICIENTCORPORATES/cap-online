"""
Stage 3 of the CS/CMA ToC-extraction pipeline: SPLIT each consolidated
subject PDF into one PDF per chapter, using the page ranges from
telegram/source-docs/CS_CMA_Chapter_Catalog.xlsx (built by
scan_cs_cma_toc.py + build_cs_cma_catalog.py -- see
_claude/skills/SKILL-cs-cma-toc-pipeline.md for the full pipeline).

SAFETY (read before running):
  - Defaults to a DRY RUN: prints exactly what it WOULD do (source file,
    page range, output filename) without writing a single file. You must
    pass --execute to actually split and write PDFs.
  - On --execute, WIPES and recreates --output-dir from scratch, so a
    re-run is always a clean, reproducible rebuild from whatever the
    Excel catalog currently says -- same idempotent-pipeline pattern as
    telegram/tools/build_study_bot_catalog.py for the CA bot. It does
    NOT touch the source PDFs under telegram/assets/{CS *, CMA *}/,
    which stay read-only inputs.
  - Every row is validated independently before anything is written --
    a bad row (no page range, out-of-bounds range, missing source file)
    is skipped and logged, never a crash that aborts the whole batch.
  - Re-checks FinalPDFFileName uniqueness itself rather than trusting
    Stage 2's own check, in case the Excel was hand-edited since.

Run (safe, no writes):   python telegram/tools/split_cs_cma_pdfs.py
Run (writes real files): python telegram/tools/split_cs_cma_pdfs.py --execute
Custom output folder:    python telegram/tools/split_cs_cma_pdfs.py --execute --output-dir path\to\folder
"""

import argparse
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cs_cma_common import ASSETS_ROOT, REPO_ROOT

EXCEL_PATH = REPO_ROOT / "telegram" / "source-docs" / "CS_CMA_Chapter_Catalog.xlsx"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "telegram" / "assets" / "cs_cma_flat"
SHEET_NAME = "Chapter Catalog"


def load_catalog_rows(excel_path: Path):
    if not excel_path.exists():
        raise SystemExit(
            f"Catalog not found: {excel_path}\n"
            f"Run build_cs_cma_catalog.py first (after scan_cs_cma_toc.py)."
        )
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet {SHEET_NAME!r} not found in {excel_path} -- found {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    required = {"SourceFolder", "SourceFile", "PDFPageStart_1indexed", "PDFPageEnd_1indexed",
                "FinalPDFFileName", "ChapterNo", "ChapterName"}
    missing = required - set(headers)
    if missing:
        raise SystemExit(f"Excel is missing expected column(s): {missing}")
    rows = [dict(zip(headers, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return rows


def validate_rows(rows):
    """Returns (usable_rows, skipped_rows). Never raises -- every problem
    with a single row is reported and that row is set aside, so one bad
    row can never block the rest of the batch."""
    usable, skipped = [], []
    seen_filenames = {}  # FinalPDFFileName -> "folder/file ch<N>" that claimed it first

    for row in rows:
        loc = f"{row.get('SourceFolder')}/{row.get('SourceFile')} ch{row.get('ChapterNo')}"
        reason = None

        start, end = row.get("PDFPageStart_1indexed"), row.get("PDFPageEnd_1indexed")
        if start is None or end is None:
            reason = f"no page range parsed (Notes: {row.get('Notes') or 'none'})"
        elif not isinstance(start, (int, float)) or not isinstance(end, (int, float)):
            reason = f"non-numeric page range ({start!r}, {end!r})"
        else:
            start_i, end_i = int(start), int(end)
            if start_i < 1:
                reason = f"start page {start_i} < 1"
            elif start_i > end_i:
                reason = f"start page {start_i} > end page {end_i}"

        if reason is None:
            fname = row.get("FinalPDFFileName")
            if not fname:
                reason = "no FinalPDFFileName"
            elif fname in seen_filenames:
                reason = f"duplicate FinalPDFFileName (already claimed by {seen_filenames[fname]})"

        if reason:
            skipped.append((row, reason))
        else:
            seen_filenames[row["FinalPDFFileName"]] = loc
            usable.append(row)

    return usable, skipped


def group_by_source(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[(row["SourceFolder"], row["SourceFile"])].append(row)
    return grouped


def split_all(rows, output_dir: Path, execute: bool):
    usable, skipped = validate_rows(rows)

    print(f"Catalog rows: {len(rows)}   usable: {len(usable)}   skipped up front: {len(skipped)}")
    if skipped:
        print("\nSkipped before opening any PDF (logged, batch continues):")
        for row, reason in skipped:
            print(f"  {row.get('SourceFolder')}/{row.get('SourceFile')} ch{row.get('ChapterNo')} "
                  f"{row.get('ChapterName')!r}: {reason}")

    grouped = group_by_source(usable)

    if execute:
        if output_dir.exists():
            shutil.rmtree(output_dir)
        output_dir.mkdir(parents=True)
        print(f"\n(--execute) output dir wiped and recreated: {output_dir}")
    else:
        print(f"\n(dry run) nothing will be written; target would be: {output_dir}")

    total_written = 0
    total_pages = 0
    errors = []

    for (folder, filename), file_rows in sorted(grouped.items()):
        src_path = ASSETS_ROOT / folder / filename
        if not src_path.exists():
            for row in file_rows:
                errors.append((row, f"source file not found: {src_path}"))
            continue

        try:
            reader = PdfReader(str(src_path))
            n_pages = len(reader.pages)
        except Exception as e:
            for row in file_rows:
                errors.append((row, f"could not open source PDF: {e!r}"))
            continue

        print(f"\n{folder}/{filename} ({n_pages} pdf pages, {len(file_rows)} chapters):")

        for row in sorted(file_rows, key=lambda r: r["ChapterNo"]):
            start_1idx, end_1idx = int(row["PDFPageStart_1indexed"]), int(row["PDFPageEnd_1indexed"])
            start_idx, end_idx = start_1idx - 1, end_1idx - 1  # pypdf reader.pages[] is 0-indexed

            if start_idx < 0 or end_idx >= n_pages:
                errors.append((row, f"page range {start_1idx}-{end_1idx} (1-indexed) is out of "
                                     f"bounds for this {n_pages}-page source PDF"))
                continue

            out_name = row["FinalPDFFileName"]
            out_path = output_dir / out_name
            n_pages_this_chapter = end_idx - start_idx + 1

            if execute:
                writer = PdfWriter()
                for idx in range(start_idx, end_idx + 1):
                    writer.add_page(reader.pages[idx])
                with open(out_path, "wb") as f:
                    writer.write(f)

            total_pages += n_pages_this_chapter
            total_written += 1
            verb = "wrote" if execute else "would write"
            print(f"  ch{row['ChapterNo']:>3}  pdf pg {start_1idx}-{end_1idx} "
                  f"({n_pages_this_chapter:>3} pages)  {verb}  {out_name}")

    print(f"\n{'Wrote' if execute else 'Would write'} {total_written} chapter PDFs "
          f"({total_pages} total pages) {'into' if execute else 'to'} {output_dir}")

    all_problems = skipped + errors
    if all_problems:
        print(f"\n{len(all_problems)} row(s) NOT split -- review before trusting this as complete:")
        for row, reason in all_problems:
            print(f"  {row.get('SourceFolder')}/{row.get('SourceFile')} ch{row.get('ChapterNo')} "
                  f"{row.get('ChapterName')!r}: {reason}")

    if not execute:
        print("\nThis was a DRY RUN -- no files were written or deleted. "
              "Re-run with --execute once you're happy with the plan above.")

    return total_written, all_problems


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--execute", action="store_true",
                         help="Actually split and write PDFs (wipes --output-dir first). "
                              "Without this flag, performs a dry run only -- no files touched.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR,
                         help=f"Where to write split chapter PDFs (default: {DEFAULT_OUTPUT_DIR})")
    args = parser.parse_args()

    rows = load_catalog_rows(EXCEL_PATH)
    _, problems = split_all(rows, args.output_dir, args.execute)
    sys.exit(1 if problems and args.execute else 0)


if __name__ == "__main__":
    main()
