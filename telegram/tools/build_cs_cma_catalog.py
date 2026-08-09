"""
Stage 2 of the CS/CMA ToC-extraction pipeline: turn the verified chapter
list from scan_cs_cma_toc.py (telegram/tools/_reports/cs_cma_toc_scan.json)
into the Excel catalog Pranav asked for.

This stage does NOT split any PDFs yet -- it only specifies, per chapter:
  - the printed page range (as a human reads it in the book)
  - the actual PDF page range (1-indexed -- subtract 1 for pypdf's
    0-indexed `reader.pages[]` access when a future splitting script
    reads this catalog)
  - a full chapter name and a short, Telegram-button-width chapter name
  - the planned final PDF filename for that chapter, once split

Two known source-data issues are corrected here via an explicit,
documented override table (never silently): two genuine typos in
ICMAI's own printed ToC ("Operatinal" / "Diefferent"), and one ICSI
lesson that its own book states was merged into an earlier lesson and
has no standalone content -- see SKILL-cs-cma-toc-pipeline.md §4.

Run: python telegram/tools/build_cs_cma_catalog.py
(after telegram/tools/scan_cs_cma_toc.py has produced a fresh
_reports/cs_cma_toc_scan.json)
"""

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cs_cma_common import REPO_ROOT, normalize_ws

SCAN_REPORT = Path(__file__).resolve().parent / "_reports" / "cs_cma_toc_scan.json"
EXCEL_PATH = REPO_ROOT / "telegram" / "source-docs" / "CS_CMA_Chapter_Catalog.xlsx"

# (folder, filename, chapter_number) -> corrected title.
# Every entry here was confirmed by reading the chapter's actual opening
# page against ICMAI/ICSI's own printed ToC -- see the scan report's
# verify_score for why each was flagged (both scored <45 against the
# ToC-parsed title before this correction).
TITLE_OVERRIDES = {
    ("CMA Final", "Cost and Mgmt Audit.pdf", 15): "Operational Audit and Internal Audit under Companies Act, 2013",
    ("CMA Final", "Cost and Mgmt Audit.pdf", 16): "Audit of Different Service Organisations",
}

# (folder, filename, chapter_number) -> reason it has no standalone page
# range and should be listed with a note instead of a real page range.
# Confirmed by reading the book's own ToC text for that lesson.
MERGED_NO_CONTENT = {
    ("CS Prof", "ENVIRONMENTAL_SOCIAL_AND_GOVERNANCE_ESG_PRINCIPLES_PRACTICE_14082025.pdf", 6):
        "This lesson number was merged into Lesson 3 (Board Effectiveness / Building Better Boards) "
        "by ICSI itself -- the book's own ToC states so explicitly. No standalone page range exists.",
}


def short_title(title: str, max_len: int = 35) -> str:
    t = re.sub(r"\bAccounting Standard\b", "AS", title, flags=re.IGNORECASE)
    t = re.sub(r"\bIndian Accounting Standard\b", "IndAS", t, flags=re.IGNORECASE)
    t = t.replace("&", "and")
    t = re.sub(r"[^A-Za-z0-9 ]", " ", t)
    words = [w for w in t.split() if w]
    out = ""
    for w in words:
        piece = w[0].upper() + w[1:]
        if len(out) + len(piece) > max_len:
            break
        out += piece
    return out or "Chapter"


def short_subject(subject: str, max_len: int = 20) -> str:
    return short_title(subject, max_len)


LEVEL_CODE = {
    "Executive": "Exec", "Professional": "Prof", "CSEET": "EET",
    "Final": "Final", "Foundation": "Found", "Intermediate": "Inter",
}


def build_rows():
    data = json.loads(SCAN_REPORT.read_text(encoding="utf-8"))
    rows = []
    for rec in data:
        folder, filename = rec["folder"], rec["filename"]
        course, level, subject = rec["Course"], rec["Level"], rec["Subject"]
        group, paper_no = rec.get("Group"), rec.get("PaperNo")
        subj_short = short_subject(subject)
        level_code = LEVEL_CODE.get(level, level)
        prefix = f"{course}{level_code}-{subj_short}"
        publisher = rec["publisher"]
        code_prefix = "M" if publisher == "icmai" else "L"

        for ch in rec["chapters"]:
            num = ch["number"]
            override_key = (folder, filename, num)
            title = TITLE_OVERRIDES.get(override_key, ch["title"])
            title = normalize_ws(title)
            merged_note = MERGED_NO_CONTENT.get(override_key)

            title_was_corrected = override_key in TITLE_OVERRIDES
            if title_was_corrected:
                note = (
                    "Title corrected -- ICMAI's own printed ToC has a typo here "
                    f"({ch['title']!r}); confirmed against the chapter's actual opening page. "
                    "Page range is correct."
                )
            elif merged_note:
                note = merged_note
            elif ch.get("verify_score", 100) < 70:
                note = "LOW CONFIDENCE -- verify manually before splitting"
            else:
                note = ""

            code = f"{code_prefix}{num}"
            row = {
                "Course": course,
                "Level": level,
                "Group": group or "",
                "PaperNo": paper_no,
                "Subject": subject,
                "ChapterCode": code,
                "ChapterNo": num,
                "ChapterName": title,
                "ShortChapterName": short_title(title),
                "PrintedPageStart": ch.get("printed_start"),
                "PrintedPageEnd": ch.get("printed_end"),
                "PDFPageStart_1indexed": (ch["pdf_start"] + 1) if ch.get("pdf_start") is not None else None,
                "PDFPageEnd_1indexed": (ch["pdf_end"] + 1) if ch.get("pdf_end") is not None else None,
                "VerifyScore": ch.get("verify_score"),
                "SourceFolder": folder,
                "SourceFile": filename,
                "FinalPDFFileName": f"{prefix}_{code}_{short_title(title)}.pdf",
                "Notes": note,
            }
            rows.append(row)
    return rows


def write_excel(rows):
    wb = Workbook()
    readme = wb.active
    readme.title = "Read Me First"
    readme.column_dimensions["A"].width = 105
    readme["A1"] = "CS / CMA Chapter Catalog (auto-generated)"
    readme["A1"].font = Font(bold=True, size=13)
    lines = [
        "",
        "This file is GENERATED by telegram/tools/build_cs_cma_catalog.py from",
        "telegram/tools/scan_cs_cma_toc.py's output. Do not hand-edit rows here --",
        "fix the source (TITLE_OVERRIDES / MERGED_NO_CONTENT in that script, or the",
        "regex parsers in scan_cs_cma_toc.py) and re-run both scripts.",
        "",
        "Unlike the CA Study Hub bot, CS and CMA study material ships as ONE",
        "consolidated PDF per subject (no per-chapter PDFs from the institute).",
        "This catalog specifies where each chapter starts/ends WITHIN that PDF,",
        "derived from the PDF's own printed Table of Contents and verified by",
        "content-matching against the actual page text -- see VerifyScore.",
        "",
        "PrintedPageStart/End = the page number as printed in the book.",
        "PDFPageStart/End_1indexed = the actual PDF page (as any PDF viewer",
        "  shows it, 1-indexed). If you write a splitting script with pypdf,",
        "  subtract 1 from both to get pypdf's 0-indexed reader.pages[] index.",
        "SourceFile = which consolidated PDF (under telegram/assets/<SourceFolder>/)",
        "  this chapter's pages come from -- nothing has been split out yet.",
        "FinalPDFFileName = the planned filename for this chapter once it IS split",
        "  into its own flat PDF (not created by this script).",
        "VerifyScore = how confidently the computed PDF page range was confirmed",
        "  against the chapter's actual opening-page text (0-100). Anything below",
        "  70 is flagged in the Notes column -- read those before trusting the row.",
        "",
        "See _claude/skills/SKILL-cs-cma-toc-pipeline.md for the full pipeline.",
    ]
    for i, line in enumerate(lines, start=2):
        readme[f"A{i}"] = line

    ws = wb.create_sheet("Chapter Catalog")
    headers = [
        "Course", "Level", "Group", "PaperNo", "Subject", "ChapterCode", "ChapterNo",
        "ChapterName", "ShortChapterName", "PrintedPageStart", "PrintedPageEnd",
        "PDFPageStart_1indexed", "PDFPageEnd_1indexed", "VerifyScore",
        "SourceFolder", "SourceFile", "FinalPDFFileName", "Notes",
    ]
    widths = [6, 12, 9, 8, 42, 10, 9, 55, 35, 14, 12, 18, 16, 11, 20, 45, 65, 55]
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"

    low_conf_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    for row_idx, r in enumerate(rows, start=2):
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=r[h])
            if r["Notes"]:
                cell.fill = low_conf_fill

    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)


def main():
    rows = build_rows()
    write_excel(rows)
    flagged = [r for r in rows if r["Notes"]]
    print(f"{len(rows)} chapter rows written to {EXCEL_PATH}")
    print(f"Rows with a Note (merged-lesson / low-confidence): {len(flagged)}")
    for r in flagged:
        print(f"  {r['SourceFolder']}/{r['SourceFile']} ch{r['ChapterNo']} {r['ChapterName']!r}: {r['Notes']}")

    # sanity: FinalPDFFileName must be globally unique
    names = [r["FinalPDFFileName"] for r in rows]
    dupes = {n for n in names if names.count(n) > 1}
    if dupes:
        print(f"\n!!! {len(dupes)} DUPLICATE planned filenames found:")
        for d in dupes:
            print(" ", d)
    else:
        print("\nOK: all planned FinalPDFFileName values are globally unique.")


if __name__ == "__main__":
    main()
