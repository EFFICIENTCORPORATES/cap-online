"""
Stage 2 of the Study Bot catalog pipeline: BUILD the flat, bot-ready PDF
folder + the "File Mapping" Excel catalog from the nested ICAI source
tree, after it has been verified by scan_study_bot_source.py.

Reads (read-only):
    telegram/assets/study_bot/<course-slug>/Module N/*.pdf

Writes:
    telegram/assets/study_bot_flat/*.pdf                          (flat, renamed copies)
    telegram/source-docs/1Lavya_Study_Hub_File_Mapping.xlsx        (generated catalog)

This script is idempotent: it always regenerates both outputs from
scratch (wipes and rewrites the flat folder + the Excel file), so it is
always safe to re-run after new PDFs are added to the source tree, or
after editing COURSE_META / TITLE_OVERRIDES below.

COURSE_META (course-slug -> Course/Level/Subject/Session) was built by
reading each course's real "Initial Pages" cover PDF, not guessed --
see _reports/_tmp_initial_pages_all.txt for the source text.

TITLE_OVERRIDES corrects the handful of filenames on disk that were
wrong/unhelpful (a "Corrigendum" with no M_C_U prefix, two chapters
literally named "Untitled", two Statistics units filed under C0/U0) --
each was confirmed by reading the actual PDF's first page, not guessed.
See _claude/skills/SKILL-study-bot-catalog-pipeline.md for the full
writeup of why each override exists.

Run: python telegram/tools/build_study_bot_catalog.py
"""

import re
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader

REPO_ROOT = Path(__file__).resolve().parents[2]
SOURCE_ROOT = REPO_ROOT / "telegram" / "assets" / "study_bot"
FLAT_ROOT = REPO_ROOT / "telegram" / "assets" / "study_bot_flat"
EXCEL_PATH = REPO_ROOT / "telegram" / "source-docs" / "1Lavya_Study_Hub_File_Mapping.xlsx"

FNAME_RE = re.compile(r"^M(?P<module>\d+)_C(?P<chapter>\d+)_U(?P<unit>\d+)_\s*(?P<title>.+)$")
CHAPTER_IN_TITLE_RE = re.compile(r"^Chapter\s+(\d+)\s*(.*)$", re.IGNORECASE)

# course-slug -> (Course, Level, Subject (full, verified from cover page),
#                 SubjectShort (for filenames), PaperLabel, SessionShort)
COURSE_META = {
    "ca-final-advaudit-may26":     ("CA", "Final", "Advanced Auditing, Assurance & Professional Ethics", "AdvAudit", "Paper 3", "May26"),
    "ca-final-afm-may26":          ("CA", "Final", "Advanced Financial Management", "AFM", "Paper 2", "May26"),
    "ca-final-directtax-nov26":    ("CA", "Final", "Direct Tax Laws & International Taxation", "DirectTax", "Paper 4", "Nov26"),
    "ca-final-finreporting-may26": ("CA", "Final", "Financial Reporting", "FinReporting", "Paper 1", "May26"),
    "ca-final-indirecttax-nov26":  ("CA", "Final", "Indirect Tax Laws (Part I: GST)", "IndirectTax", "Paper 5", "Nov26"),
    "ca-found-accounts-may26":     ("CA", "Foundation", "Accounting", "Accounts", "Paper 1", "May26"),
    "ca-found-businessslaw-may26": ("CA", "Foundation", "Business Laws", "BusinessLaw", "Paper 2", "May26"),
    "ca-found-eco-may26":          ("CA", "Foundation", "Business Economics", "Eco", "Paper 4", "May26"),
    "ca-found-quants-may26":       ("CA", "Foundation", "Quantitative Aptitude", "Quants", "Paper 3", "May26"),
    "ca-inter-advacc-may26":       ("CA", "Inter", "Advanced Accounting", "AdvAcc", "Paper 1", "May26"),
    "ca-inter-audit-may27":        ("CA", "Inter", "Auditing & Ethics", "Audit", "Paper 5", "May27"),
    "ca-inter-costacc-may27":      ("CA", "Inter", "Cost and Management Accounting", "CostAcc", "Paper 4", "May27"),
    "ca-inter-fm-may26":           ("CA", "Inter", "Financial Management", "FM", "Paper 6A", "May26"),
    # GST has TWO simultaneously-live ICAI editions as of 2026-08-18 (the
    # May26/Sep26/Jan27 attempt cycle and the May27/Sep27/Jan28 cycle both
    # remain relevant to real, currently-enrolled students) -- kept as two
    # distinct Subject strings (not one "GST" entry) specifically so
    # populate_course_catalog.py's per-subject chapter_no grouping never
    # sees two editions' identically-numbered chapters as one fabricated
    # "multi-unit chapter" collision. See telegram/COURSE-CATALOG.md.
    "ca-inter-gsttax-may26":       ("CA", "Inter", "Taxation - Goods and Services Tax (May 2026/Sep 2026/Jan 2027 Attempt)", "GST", "Paper 3B", "May26"),
    "ca-inter-gsttax-may27":       ("CA", "Inter", "Taxation - Goods and Services Tax (May 2027/Sep 2027/Jan 2028 Attempt)", "GST", "Paper 3B", "May27"),
    "ca-inter-incometax-may27":    ("CA", "Inter", "Taxation - Income-tax Law", "IncomeTax", "Paper 3A", "May27"),
    "ca-inter-law-may27":          ("CA", "Inter", "Corporate and Other Laws", "Law", "Paper 2", "May27"),
    "ca-sm-inter-p6b-may2027":     ("CA", "Inter", "Strategic Management", "SM", "Paper 6B", "May27"),
}

# (course_slug, on-disk filename) -> corrected (module, chapter, unit, title)
# Every override below was confirmed by reading the PDF's actual first page
# with pypdf -- see _reports/_tmp_untitled_full.txt and
# _reports/_tmp_quants_full.txt for the raw extracted text that justifies
# each correction.
TITLE_OVERRIDES = {
    ("ca-found-accounts-may26", "ca-found-accounts-may26-corrigendum.pdf"):
        (1, 0, 9, "Corrigendum"),
    ("ca-final-advaudit-may26", "M2_C14_U1_ Untitled.pdf"):
        (2, 14, 1, "Special Features of Audit of Banks"),
    ("ca-final-advaudit-may26", "M2_C14_U2_ Untitled.pdf"):
        (2, 14, 2, "Special Features of Audit of Non-Banking Financial Companies (NBFCs)"),
    ("ca-found-quants-may26", "M1_C0_U0_ Unit I Statistical Description of Data.pdf"):
        (1, 13, 1, "Statistical Description of Data"),
    ("ca-found-quants-may26", "M1_C0_U0_ Unit I Measures of Central Tendency.pdf"):
        (1, 14, 1, "Measures of Central Tendency and Dispersion"),
}


def parse_filename(course_slug: str, pdf_path: Path):
    stem = pdf_path.stem
    key = (course_slug, pdf_path.name)
    if key in TITLE_OVERRIDES:
        module, chapter, unit, title = TITLE_OVERRIDES[key]
        return module, chapter, unit, title

    m = FNAME_RE.match(stem)
    if not m:
        raise ValueError(f"Unparseable filename with no override: {course_slug}/{pdf_path.name}")

    module = int(m.group("module"))
    chapter = int(m.group("chapter"))
    unit = int(m.group("unit"))
    title = m.group("title").strip()

    # ca-found-quants-style files: M_C_U prefix is uninformative (0/0), but
    # the title itself states "Chapter N <real title>" -- promote it.
    if chapter == 0 and module in (0, 1):
        cm = CHAPTER_IN_TITLE_RE.match(title)
        if cm:
            chapter = int(cm.group(1))
            title = cm.group(2).strip() or title

    return module, chapter, unit, title


def short_title(title: str, max_len: int = 35) -> str:
    t = re.sub(r"^(Chapter\s+\d+|Unit\s+[IVXLCM]+)\s*[:\-]?\s*", "", title, flags=re.IGNORECASE)
    t = re.sub(r"\bAccounting Standard\b", "AS", t, flags=re.IGNORECASE)
    t = re.sub(r"\bIndian Accounting Standard\b", "IndAS", t, flags=re.IGNORECASE)
    t = t.replace("&", "and")
    t = re.sub(r"[^A-Za-z0-9 ]", " ", t)
    words = [w for w in t.split() if w]
    out = ""
    for w in words:
        piece = w[0].upper() + w[1:]
        if len(out) + len(piece) > max_len:
            break
        out += piece
    return out or "Chapter"


def build_keywords(subject_full: str, chapter_name: str) -> str:
    kw = {subject_full.lower(), chapter_name.lower()}
    m = re.search(r"\bAS\s?(\d+)\b", chapter_name) or re.search(r"Accounting Standard\s+(\d+)", chapter_name, re.IGNORECASE)
    if m:
        kw.add(f"as {m.group(1)}")
    m2 = re.search(r"Ind\s?AS\s?(\d+)", chapter_name, re.IGNORECASE)
    if m2:
        kw.add(f"ind as {m2.group(1)}")
    return ", ".join(sorted(kw))


def main():
    if FLAT_ROOT.exists():
        shutil.rmtree(FLAT_ROOT)
    FLAT_ROOT.mkdir(parents=True)

    rows = []
    seen_filenames = set()
    missing_meta = set()

    course_dirs = sorted(p for p in SOURCE_ROOT.iterdir() if p.is_dir())
    for course_dir in course_dirs:
        course_slug = course_dir.name
        if course_slug not in COURSE_META:
            missing_meta.add(course_slug)
            continue
        course, level, subject_full, subject_short, paper_label, session = COURSE_META[course_slug]

        for pdf_path in sorted(course_dir.rglob("*.pdf")):
            module, chapter, unit, title = parse_filename(course_slug, pdf_path)

            prefix = f"CA{level}-{subject_short}-{session}"
            code = f"M{module}-C{chapter}-U{unit}"
            new_name = f"{prefix}_{code}_{short_title(title)}.pdf"

            if new_name in seen_filenames:
                # deterministic disambiguation, should not normally trigger
                n = 2
                base = new_name[:-4]
                while f"{base}-{n}.pdf" in seen_filenames:
                    n += 1
                new_name = f"{base}-{n}.pdf"
            seen_filenames.add(new_name)

            shutil.copy2(pdf_path, FLAT_ROOT / new_name)

            rows.append({
                "FileName": new_name,
                "Course": course,
                "Level": level,
                "Subject": subject_full,
                "ModuleOrGroup": f"Module {module}",
                "ChapterNo": chapter,
                "ChapterName": title,
                "Keywords": build_keywords(subject_full, title),
                "_sort": (course_slug, module, chapter, unit),
            })

    if missing_meta:
        raise SystemExit(f"COURSE_META is missing entries for: {sorted(missing_meta)}")

    rows.sort(key=lambda r: r["_sort"])
    for r in rows:
        del r["_sort"]

    write_excel(rows)

    print(f"PDFs copied to flat folder: {len(rows)}  ->  {FLAT_ROOT}")
    print(f"Excel catalog written: {EXCEL_PATH}  ({len(rows)} rows)")
    src_count = sum(len(list(d.rglob('*.pdf'))) for d in course_dirs)
    assert src_count == len(rows), f"Source has {src_count} PDFs but only {len(rows)} rows were produced"
    assert len(seen_filenames) == len(rows), "Filename collision slipped through"
    print("OK: source PDF count == flat PDF count == Excel row count, all filenames unique.")


def write_excel(rows):
    wb = Workbook()

    readme = wb.active
    readme.title = "Read Me First"
    readme.column_dimensions["A"].width = 100
    bold = Font(bold=True, size=13)
    readme["A1"] = "1Lavya Study Hub — File Mapping (auto-generated)"
    readme["A1"].font = bold
    readme_lines = [
        "",
        "This file is GENERATED by telegram/tools/build_study_bot_catalog.py.",
        "Do not hand-edit rows here -- edit the source PDFs / TITLE_OVERRIDES in",
        "that script and re-run it, so this catalog and the flat PDF folder never",
        "drift apart.",
        "",
        "PDFs live in: telegram/assets/study_bot_flat/ (flat, one folder, no subfolders).",
        "Source of truth for content: telegram/assets/study_bot/ (nested, as ICAI",
        "shipped it) -- kept as the archival source; not read by the bot directly.",
        "",
        "Column A (FileName) exactly matches the real file name on disk, including .pdf.",
        "See _claude/skills/SKILL-study-bot-catalog-pipeline.md for the full pipeline.",
    ]
    for i, line in enumerate(readme_lines, start=2):
        readme[f"A{i}"] = line

    ws = wb.create_sheet("File Mapping")
    headers = ["FileName", "Course", "Level", "Subject", "ModuleOrGroup", "ChapterNo", "ChapterName", "Keywords"]
    widths = [70, 8, 12, 40, 14, 10, 55, 45]
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    for col_idx, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"

    for row_idx, r in enumerate(rows, start=2):
        for col_idx, h in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=r[h])

    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)


if __name__ == "__main__":
    main()
